import streamlit as st
import io
import os
import re
import csv
import zipfile
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ૧. પેજ સેટઅપ
st.set_page_config(
    page_title="Bank Report to Excel Converter",
    page_icon="🏦",
    layout="centered"
)

st.markdown("""
    <style>
    .main-title {
        text-align: center;
        color: #1F4E78;
        font-size: 26px;
        font-weight: bold;
        margin-bottom: 5px;
    }
    .sub-title {
        text-align: center;
        color: #555555;
        font-size: 15px;
        margin-bottom: 20px;
    }
    .stDownloadButton button {
        width: 100%;
        background-color: #1F4E78;
        color: white;
        font-weight: bold;
        height: 50px;
        border-radius: 8px;
    }
    </style>
""", unsafe_allow_html=True)

st.markdown('<div class="main-title">🏦 Bank Reports to Excel Converter</div>', unsafe_allow_html=True)
st.markdown('<div class="sub-title">ટેક્સ્ટ રિપોર્ટ ફાઇલોને ૧-ક્લિકમાં પ્રોપર એક્સેલમાં કન્વર્ટ કરો</div>', unsafe_allow_html=True)

# ૨. ફિલ્ટર અને ક્લીનિંગ લોજિક
IGNORE_PATTERNS = [
    "REPORT ID:", "PROC DATE:", "RUN DATE:", "BRANCH :", "BRANCH NO.", "PAGE NO",
    "PRODUCT TOTAL", "ACCOUNT TYPE TOTAL", "OVER DRAWN TOTAL", "NO OF ACCOUNTS",
    "SUB TOTAL", "GRAND TOTAL", "BALANCE FORWARD", "BHAVNAGAR DISTRICT",
    "TOTAL FOR PRODUCT", "TOTAL NO OF", "PRODUCT-WISE TOTAL", "====>",
    "GL CLASS CODE", "GL-CLASS-CODE", "AREA:"
]

def clean_dr_amount(val):
    val_str = str(val).strip()
    if not val_str:
        return ""
    if re.search(r'\bDr\b', val_str, re.IGNORECASE) or val_str.endswith("Dr") or val_str.endswith("dr"):
        num_clean = re.sub(r'(?i)\s*dr\s*', '', val_str).strip()
        return f"-{num_clean}" if not num_clean.startswith("-") else num_clean
    elif re.search(r'\bCr\b', val_str, re.IGNORECASE):
        return re.sub(r'(?i)\s*cr\s*', '', val_str).strip()
    return val_str

def parse_deposits_balance_report(lines):
    pattern = re.compile(
        r'^\s*(\d{11}-\d|\d{8,16})\s+'
        r'(\S+(?:\s\S+)*)\s{2,}'
        r'(.+?)\s{2,}'
        r'([\d,]+\.\d{2})\s+'
        r'([\d,]+\.\d{2})\s+'
        r'([\d,]+\.\d{2}(?:\s*Dr)?)\s+'
        r'([\d,]+\.\d{2})\s*'
        r'(?:\s+(\d+[DMY]))?'
        r'\s+([\d,]+\.\d{2})'
        r'\s+([A-Z]+)'
        r'\s+([YN])\s*$',
        re.IGNORECASE
    )
    rows = []
    for l in lines:
        m = pattern.match(l)
        if m:
            rows.append({
                "ACCOUNT NUMBER": m.group(1).strip(),
                "ACCOUNT TYPE (DESCRIPTION)": m.group(2).strip(),
                "CUSTOMER NAME": m.group(3).strip(),
                "AVAILABLE BALANCE": m.group(4).strip(),
                "UNCLEARED BALANCE": m.group(5).strip(),
                "CURRENT BALANCE": clean_dr_amount(m.group(6)),
                "LIMIT": m.group(7).strip(),
                "TERM": m.group(8).strip() if m.group(8) else "",
                "INT-RATE": m.group(9).strip(),
                "STATUS": m.group(10).strip(),
                "JOINT-HOLD-FLAG": m.group(11).strip()
            })
    return pd.DataFrame(rows) if rows else None

def parse_pipe_delimited(lines):
    header_cols = []
    data_rows = []
    for l in lines:
        st_ = l.strip()
        if not st_ or re.match(r'^-{10,}', st_) or st_.startswith('^[') or st_.startswith('\x0c') or any(kw in st_.upper() for kw in IGNORE_PATTERNS):
            continue
        if '|' in st_:
            parts = [clean_dr_amount(p.strip()) for p in st_.split('|')]
            if parts and parts[0] == "": parts = parts[1:]
            if parts and parts[-1] == "": parts = parts[:-1]
            if not header_cols and any(w in st_.upper() for w in ["NAME", "ACCOUNT", "AMOUNT", "LIMIT", "DATE", "PRODUCT"]):
                header_cols = parts
            elif header_cols and len(parts) >= 2:
                if len(parts) < len(header_cols): parts.extend([""] * (len(header_cols) - len(parts)))
                data_rows.append(parts[:len(header_cols)])
    return pd.DataFrame(data_rows, columns=header_cols) if header_cols and data_rows else None

def parse_general_cbs_report(lines):
    h_idx = -1
    for i, l in enumerate(lines):
        if re.match(r'^\s*-{15,}', l) and i + 2 < len(lines) and re.match(r'^\s*-{15,}', lines[i+2]):
            h_idx = i + 1; break
    if h_idx == -1: return None
    cols = [m.group(1).strip() for m in re.finditer(r'(\S+(?:\s(?!\s)\S+)*)', lines[h_idx])]
    if len(cols) < 2: return None
    rows = []
    for l in lines[h_idx+2:]:
        st_ = l.strip()
        if not st_ or re.match(r'^-{10,}', st_) or st_.startswith('^[') or st_.startswith('\x0c') or any(kw in st_.upper() for kw in IGNORE_PATTERNS) or "ACCOUNT NUMBER" in st_:
            continue
        parts = [p.strip() for p in re.split(r'\s{2,}', st_) if p.strip() != ""]
        if len(parts) > 1:
            rows.append({cols[k]: clean_dr_amount(parts[k]) if k < len(parts) else "" for k in range(len(cols))})
    return pd.DataFrame(rows) if rows else None

def parse_csv_generic(content):
    try: delim = csv.Sniffer().sniff(content[:4096], delimiters=[',', '\t', '|', ';']).delimiter
    except: delim = ',' if ',' in content else '\t'
    df = pd.read_csv(io.StringIO(content), sep=delim, engine="python", on_bad_lines='skip', encoding_errors='ignore')
    for col in df.columns:
        if df[col].dtype == object: df[col] = df[col].apply(clean_dr_amount)
    return df

def convert_to_excel(df):
    output = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.append(list(df.columns))
    h_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    h_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    border = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'), top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))
    
    for cell in ws[1]:
        cell.fill = h_fill; cell.font = h_font; cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in df.itertuples(index=False): 
        ws.append(list(row))
    for r in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in r:
            cell.font = data_font; cell.border = border
            val_s = str(cell.value or "").strip()
            cell.alignment = Alignment(horizontal="right") if re.match(r'^-?[\d,]+(\.\d+)?$', val_s) else Alignment(horizontal="left")
    
    # Auto Column Width (સુધારેલ ભાગ: cell.value)
    for col in ws.columns:
        max_l = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(max_l + 3, 12)
        
    ws.freeze_panes = "A2"
    wb.save(output)
    output.seek(0)
    return output.getvalue()

# ૩. યુઝર ઇન્ટરફેસ
uploaded_files = st.file_uploader(
    "ટેક્સ્ટ ફાઇલો પસંદ કરો (Select Any Files):",
    accept_multiple_files=True
)

if uploaded_files:
    total_files = len(uploaded_files)
    st.info(f"📁 પસંદ કરેલી ફાઇલોની સંખ્યા: {total_files}")
    
    if st.button("🚀 Convert to Excel (એક્સેલમાં કન્વર્ટ કરો)"):
        progress_bar = st.progress(0)
        status_text = st.empty()
        converted_files = {}
        failed_files = []
        
        for idx, file_obj in enumerate(uploaded_files):
            filename = file_obj.name
            base_name, _ = os.path.splitext(filename)
            status_text.text(f"પ્રોસેસ થઈ રહી છે: {filename} ({idx+1}/{total_files})...")
            
            try:
                content_bytes = file_obj.read()
                try: content_str = content_bytes.decode('utf-8')
                except UnicodeDecodeError: content_str = content_bytes.decode('latin1', errors='ignore')
                
                lines = content_str.splitlines()
                df = None
                if ("DEPOSITS BALANCE FILE" in content_str.upper() or "AVAILABLE BALANCE" in content_str.upper()):
                    df = parse_deposits_balance_report(lines)
                if df is None or df.empty:
                    pipe_count = sum(1 for l in lines[:50] if l.count('|') >= 3)
                    if pipe_count >= 3: df = parse_pipe_delimited(lines)
                if df is None or df.empty: df = parse_general_cbs_report(lines)
                if df is None or df.empty: df = parse_csv_generic(content_str)
                    
                if df is not None and not df.empty:
                    converted_files[f"{base_name}.xlsx"] = convert_to_excel(df)
                else:
                    failed_files.append(filename)
            except Exception as e:
                failed_files.append(f"{filename} ({str(e)})")
                
            progress_bar.progress((idx + 1) / total_files)
            
        status_text.empty()
        progress_bar.empty()
        
        st.success(f"✅ સફળતાપૂર્વક કન્વર્ટ થયેલ ફાઇલો: {len(converted_files)} / {total_files}")
        if failed_files:
            st.warning(f"⚠️ સ્કીપ થયેલ ફાઇલો: {', '.join(failed_files)}")
            
        if len(converted_files) == 1:
            only_name, only_data = list(converted_files.items())[0]
            st.download_button(label=f"📥 Download {only_name}", data=only_data, file_name=only_name, mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        elif len(converted_files) > 1:
            zip_buffer = io.BytesIO()
            with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
                for fname, fdata in converted_files.items(): zip_file.writestr(fname, fdata)
            zip_buffer.seek(0)
            st.download_button(label=f"📦 Download All Excel Files (.ZIP - {len(converted_files)} Files)", data=zip_buffer.getvalue(), file_name="Converted_Bank_Reports.zip", mime="application/zip")
